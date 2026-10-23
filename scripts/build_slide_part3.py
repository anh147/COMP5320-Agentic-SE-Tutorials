from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
SLIDES_DIR = HERE.parent / "slides"
TEMPLATE = SLIDES_DIR / "COMP5320 Slide Template.pptx"
OUTPUT = SLIDES_DIR / "slide-part3.pptx"

NAVY = RGBColor(25, 72, 119)
DEEP_NAVY = RGBColor(31, 51, 95)
RED = RGBColor(207, 31, 43)
SKY = RGBColor(224, 238, 249)
PALE_RED = RGBColor(252, 232, 234)
PALE_GREEN = RGBColor(230, 244, 237)
GREEN = RGBColor(31, 122, 78)
AMBER = RGBColor(246, 183, 66)
INK = RGBColor(34, 43, 55)
MID = RGBColor(93, 105, 119)
LIGHT = RGBColor(244, 247, 250)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(213, 220, 228)

with ZipFile(TEMPLATE) as template_zip:
    CONTENT_BACKGROUND = template_zip.read("ppt/media/image1.jpg")


def remove_slide(prs, index):
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[index]


def rect(slide, x, y, w, h, fill, radius=False, line=None, width=1):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line if line else fill
    sh.line.width = Pt(width)
    return sh


def text(slide, value, x, y, w, h, size=20, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP, margin=0.06):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rich_lines(slide, lines, x, y, w, h, size=18, color=INK, bullet=False,
               spacing=7, bullet_color=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            head, body = item
        else:
            head, body = "", item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        if bullet:
            p.text = "•  "
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = bullet_color or RED
        if head:
            r = p.add_run()
            r.text = head
            r.font.bold = True
            r.font.name = "Aptos"
            r.font.size = Pt(size)
            r.font.color.rgb = color
        r = p.add_run()
        r.text = body
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.font.color.rgb = color
    return box


def add_chrome(slide, number, section, title_text, kicker=None):
    # Use the original content-slide artwork from the supplied template instead
    # of approximating its red arrow and blue header with editable shapes.
    slide.shapes.add_picture(
        BytesIO(CONTENT_BACKGROUND), 0, 0, width=prs.slide_width, height=prs.slide_height
    )
    text(slide, section.upper(), 0.87, 0.18, 1.8, 0.25, 10, WHITE, True)
    # Keep the title fully inside the blue header. Placing dark-blue text across
    # the blue/white boundary makes its upper portion appear hidden in PowerPoint.
    text(slide, title_text, 2.55, 0.08, 9.85, 0.54, 26, WHITE, True,
         valign=MSO_ANCHOR.MIDDLE)
    if kicker:
        text(slide, kicker, 0.78, 1.31, 11.6, 0.28, 11, MID)
    text(slide, str(number), 12.42, 6.96, 0.45, 0.25, 10, MID, False, PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, heading, body=None, accent=NAVY, fill=WHITE,
             heading_size=17, body_size=15):
    rect(slide, x, y, w, h, fill, True, LINE)
    rect(slide, x, y, 0.08, h, accent)
    text(slide, heading, x + 0.24, y + 0.18, w - 0.45, 0.34, heading_size, accent, True)
    if body:
        text(slide, body, x + 0.24, y + 0.62, w - 0.45, h - 0.78, body_size, INK)


def add_pill(slide, value, x, y, w, color=NAVY, fill=SKY):
    rect(slide, x, y, w, 0.36, fill, True, fill)
    text(slide, value, x + 0.05, y + 0.035, w - 0.1, 0.25, 10, color, True, PP_ALIGN.CENTER)


def add_connector(slide, x1, y1, x2, y2, color=MID, width=2):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


prs = Presentation(TEMPLATE)
remove_slide(prs, 1)

# Title slide: retain the original photo and visual identity, update only the copy.
title_slide = prs.slides[0]
for sh in title_slide.shapes:
    if not hasattr(sh, "text_frame"):
        continue
    if "AY 25" in sh.text:
        sh.text_frame.paragraphs[0].runs[0].text = "AY 25–26"
    if "Advanced Software Engineering" in sh.text:
        sh.text_frame.clear()
        lines = [
            ("Advanced Software Engineering", 18, True),
            ("Week: ___", 14, False),
            ("Topic: Requirements Engineering & Specification — Part 3", 14, False),
            ("Sections 2.7–2.9", 14, False),
            ("Team: ___", 14, False),
        ]
        for i, (value, size, bold) in enumerate(lines):
            p = sh.text_frame.paragraphs[0] if i == 0 else sh.text_frame.add_paragraph()
            p.text = value
            p.space_after = Pt(2)
            for run in p.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = WHITE


def new_slide():
    return prs.slides.add_slide(prs.slide_layouts[0])


# Slide 2 — scope creep
s = new_slide()
add_chrome(s, 2, "2.7 Scope Creep", "Scope creep is uncontrolled change",
           "Change is normal. Unpriced, unapproved change is the risk.")
rect(s, 0.82, 1.9, 5.0, 3.85, PALE_RED, True, PALE_RED)
text(s, "Original boundary", 1.12, 2.16, 2.0, 0.68, 17, RED, True)
for i, label in enumerate(["Create tasks", "Assign owners", "Track status"]):
    add_pill(s, label, 1.15 + (i % 2) * 1.75, 2.88 + (i // 2) * 0.62, 1.55, NAVY, WHITE)
text(s, "+", 3.62, 3.04, 0.4, 0.4, 26, RED, True, PP_ALIGN.CENTER)
for i, label in enumerate(["Dark mode", "Chat", "AI summaries", "Mobile app"]):
    add_pill(s, label, 3.94, 2.48 + i * 0.61, 1.42, RED, WHITE)
text(s, "The boundary expands while time and budget stay fixed.", 1.13, 4.72, 4.25, 0.76, 17, INK, True)
for y, heading, body, color in [
    (1.90, "New requests after work starts", "A useful idea appears, but no trade-off is agreed.", RED),
    (2.96, "Ambiguous requirements", "People interpret the same story differently.", AMBER),
    (4.02, "Gold-plating by the team", "Extra features are added without approval.", NAVY),
    (5.08, "External change", "Regulation or dependencies create new work.", DEEP_NAVY),
]:
    rect(s, 6.18, y, 6.25, 0.91, WHITE, True, LINE)
    rect(s, 6.18, y, 0.08, 0.91, color)
    text(s, heading, 6.46, y + 0.10, 5.62, 0.32, 15, color, True)
    text(s, body, 6.46, y + 0.49, 5.62, 0.34, 12, INK)
text(s, "SIGNAL: “While you are there, could you also…?”", 0.92, 6.46, 11.6, 0.36, 15, RED, True, PP_ALIGN.CENTER)


# Slide 3 — change control
s = new_slide()
add_chrome(s, 3, "2.7 Scope Creep", "Make every scope change a decision",
           "A lightweight change-control loop protects value, schedule, and trust")
flow = [
    ("1", "Capture", "Write the request and why it matters."),
    ("2", "Assess", "Estimate value, effort, risk, and dependencies."),
    ("3", "Trade off", "Swap, defer, resize, or fund the change."),
    ("4", "Approve", "Name the decision owner and update the baseline."),
]
for i, (n, head, body) in enumerate(flow):
    x = 0.68 + i * 3.14
    rect(s, x, 2.05, 2.63, 2.35, WHITE, True, LINE)
    circ = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.92), Inches(1.72), Inches(0.78), Inches(0.78))
    circ.fill.solid(); circ.fill.fore_color.rgb = NAVY if i != 2 else RED; circ.line.color.rgb = circ.fill.fore_color.rgb
    text(s, n, x + 0.92, 1.86, 0.78, 0.34, 18, WHITE, True, PP_ALIGN.CENTER)
    text(s, head, x + 0.2, 2.68, 2.23, 0.4, 20, NAVY, True, PP_ALIGN.CENTER)
    text(s, body, x + 0.22, 3.25, 2.2, 0.72, 14, INK, False, PP_ALIGN.CENTER)
    if i < 3:
        text(s, "→", x + 2.68, 2.78, 0.42, 0.42, 24, MID, True, PP_ALIGN.CENTER)
rect(s, 1.05, 5.05, 11.18, 1.16, SKY, True, SKY)
text(s, "MoSCoW makes the boundary visible", 1.35, 5.28, 3.6, 0.35, 18, NAVY, True)
text(s, "Use Won’t Have (this release) to record deliberate exclusion — not forgotten work.",
     5.0, 5.23, 6.85, 0.55, 16, INK)
text(s, "RULE OF THUMB  •  No new scope without a named trade-off.", 1.2, 6.48, 10.9, 0.32, 14, RED, True, PP_ALIGN.CENTER)


# Slide 4 — acceptance criteria
s = new_slide()
add_chrome(s, 4, "2.8 Acceptance Criteria", "Acceptance criteria connect intent to evidence",
           "They define the specific, testable conditions under which a story is accepted")
add_card(s, 0.82, 2.0, 3.25, 2.88, "USER STORY", "Who needs what — and why?", NAVY, SKY, 18, 18)
add_card(s, 5.03, 2.0, 3.25, 2.88, "ACCEPTANCE CRITERIA", "What observable outcomes must hold?", RED, PALE_RED, 18, 18)
add_card(s, 9.25, 2.0, 3.25, 2.88, "TESTS", "What evidence confirms each outcome?", GREEN, PALE_GREEN, 18, 18)
text(s, "→", 4.28, 3.03, 0.5, 0.5, 26, MID, True, PP_ALIGN.CENTER)
text(s, "→", 8.5, 3.03, 0.5, 0.5, 26, MID, True, PP_ALIGN.CENTER)
rich_lines(s, [
    ("Specific — ", "one clear expected behavior"),
    ("Observable — ", "state, response, message, or side effect"),
    ("Testable — ", "pass/fail without subjective judgment"),
], 1.05, 5.42, 5.35, 1.05, 15, INK, True, 3)
rich_lines(s, [
    ("Complete — ", "happy path + errors + edge cases"),
    ("Business-facing — ", "outcomes, not implementation details"),
], 6.74, 5.42, 5.35, 0.95, 15, INK, True, 5)


# Slide 5 — Gherkin anatomy
s = new_slide()
add_chrome(s, 5, "2.8 Acceptance Criteria", "Gherkin: context → action → outcome",
           "A shared language readable by stakeholders, developers, and testers")
blocks = [
    ("GIVEN", "Initial context", "the user is authenticated\nand task 123 is unassigned", NAVY),
    ("WHEN", "Single action", "the manager assigns task 123\nto Alice", RED),
    ("THEN", "Observable outcome", "the assignee is Alice\nand a notification is queued", GREEN),
]
for i, (tag, sub, example, color) in enumerate(blocks):
    x = 0.82 + i * 4.17
    rect(s, x, 2.0, 3.62, 3.48, WHITE, True, LINE)
    rect(s, x, 2.0, 3.62, 0.7, color, True, color)
    text(s, tag, x + 0.2, 2.11, 3.22, 0.42, 20, WHITE, True, PP_ALIGN.CENTER)
    text(s, sub, x + 0.24, 2.92, 3.14, 0.4, 17, color, True, PP_ALIGN.CENTER)
    text(s, example, x + 0.3, 3.66, 3.02, 1.12, 15, INK, False, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    if i < 2:
        text(s, "→", x + 3.68, 3.45, 0.42, 0.42, 24, MID, True, PP_ALIGN.CENTER)
rect(s, 1.15, 5.93, 11.03, 0.62, LIGHT, True, LIGHT)
text(s, "Write externally visible behavior. Avoid UI clicks, classes, database tables, and algorithm choices unless they are requirements.",
     1.36, 6.06, 10.62, 0.32, 14, MID, False, PP_ALIGN.CENTER)


# Slide 6 — example scenarios
s = new_slide()
add_chrome(s, 6, "2.8 Acceptance Criteria", "One story needs more than the happy path",
           "Example story: As a project manager, I want to assign a task so ownership is clear")
scenarios = [
    ("HAPPY PATH", "Given task 123 and Alice are in the project\nWhen the manager assigns task 123 to Alice\nThen the task shows Alice as assignee\nAnd Alice is notified within 5 minutes", GREEN, PALE_GREEN),
    ("ERROR CASE", "Given Bob is not a project member\nWhen the manager assigns task 123 to Bob\nThen the request is rejected\nAnd the task remains unassigned", RED, PALE_RED),
    ("EDGE CASE", "Given task 123 is already assigned to Alice\nWhen the same assignment is repeated\nThen no duplicate notification is sent\nAnd the assignee remains Alice", NAVY, SKY),
]
for i, (head, body, color, fill) in enumerate(scenarios):
    x = 0.66 + i * 4.17
    rect(s, x, 2.03, 3.75, 3.96, fill, True, fill)
    add_pill(s, head, x + 0.83, 2.28, 2.1, color, WHITE)
    text(s, body, x + 0.31, 3.03, 3.13, 2.38, 14, INK, False, PP_ALIGN.LEFT)
text(s, "Coverage question: What could succeed, fail, or happen twice?", 1.05, 6.39, 11.2, 0.38, 16, NAVY, True, PP_ALIGN.CENTER)


# Slide 7 — acceptance criteria vs DoD
s = new_slide()
add_chrome(s, 7, "2.9 Definition of Done", "Acceptance criteria are local; the DoD is global",
           "Both must pass before a story can be called complete")
headers = ["", "Acceptance Criteria", "Definition of Done"]
xs = [0.82, 3.38, 7.82]
ws = [2.42, 4.28, 4.7]
for i, h in enumerate(headers):
    rect(s, xs[i], 2.0, ws[i], 0.7, NAVY if i != 2 else GREEN, True, NAVY if i != 2 else GREEN)
    text(s, h, xs[i] + 0.12, 2.14, ws[i] - 0.24, 0.38, 16, WHITE, True, PP_ALIGN.CENTER)
rows = [
    ("Applies to", "One specific story", "Every story / increment"),
    ("Focus", "Business behavior", "Engineering quality"),
    ("Examples", "Valid assignment; permission error", "Tests, review, security, docs"),
    ("Changes", "Story by story", "By team agreement"),
]
for r, vals in enumerate(rows):
    y = 2.84 + r * 0.79
    fills = [LIGHT, WHITE, PALE_GREEN]
    for c, val in enumerate(vals):
        rect(s, xs[c], y, ws[c], 0.66, fills[c], False, LINE)
        text(s, val, xs[c] + 0.12, y + 0.12, ws[c] - 0.24, 0.39, 14, INK, c == 0, PP_ALIGN.LEFT)
rect(s, 1.2, 6.18, 10.93, 0.58, PALE_RED, True, PALE_RED)
text(s, "DONE = story-specific outcomes pass  +  team-wide quality gate passes", 1.45, 6.31, 10.43, 0.3, 16, RED, True, PP_ALIGN.CENTER)


# Slide 8 — DoD checklist
s = new_slide()
add_chrome(s, 8, "2.9 Definition of Done", "A practical Definition of Done",
           "A shared checklist keeps “almost done” from becoming a permanent state")
items = [
    ("✓", "Behavior", "All acceptance criteria pass"),
    ("✓", "Tests", "Unit + integration tests pass; ≥80% new-code coverage"),
    ("✓", "Review", "At least one teammate approves the change"),
    ("✓", "Quality", "Linter and type checker report no errors"),
    ("✓", "Security", "Automated scan introduces no new vulnerabilities"),
    ("✓", "Release", "Docs updated; deployed and smoke-tested in staging"),
]
for i, (mark, head, body) in enumerate(items):
    col, row = i % 2, i // 2
    x, y = 0.82 + col * 6.15, 1.92 + row * 1.36
    rect(s, x, y, 5.58, 1.08, WHITE, True, LINE)
    circ = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.22), Inches(y + 0.25), Inches(0.55), Inches(0.55))
    circ.fill.solid(); circ.fill.fore_color.rgb = GREEN; circ.line.color.rgb = GREEN
    text(s, mark, x + 0.22, y + 0.32, 0.55, 0.27, 15, WHITE, True, PP_ALIGN.CENTER)
    text(s, head, x + 0.95, y + 0.14, 1.28, 0.38, 16, NAVY, True)
    text(s, body, x + 2.12, y + 0.17, 3.13, 0.62, 13, INK)
rect(s, 1.05, 6.12, 11.22, 0.65, SKY, True, SKY)
text(s, "The DoD should be ambitious enough to protect quality — and realistic enough to apply every time.",
     1.34, 6.28, 10.62, 0.3, 14, NAVY, True, PP_ALIGN.CENTER)


# Remove empty footer/date placeholders introduced by the layout while retaining slide content.
for slide in list(prs.slides)[1:]:
    for shape in list(slide.placeholders):
        if not shape.has_text_frame or not shape.text.strip():
            sp = shape._element
            sp.getparent().remove(sp)

prs.core_properties.title = "Requirements Engineering & Specification — Part 3"
prs.core_properties.subject = "Sections 2.7 Scope Creep, 2.8 Acceptance Criteria, and 2.9 Definition of Done"
prs.core_properties.author = "COMP5320 Advanced Software Engineering"
prs.core_properties.comments = "Adapted for teaching from Agentic Software Engineering, Chapter 2, sections 2.7–2.9."
prs.save(OUTPUT)
print(OUTPUT)
